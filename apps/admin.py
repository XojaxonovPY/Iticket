import datetime

import openpyxl
from django.contrib import admin, messages
from django.contrib.auth.admin import UserAdmin as BaseUserAdmin
from django.contrib.auth.models import Group
from django.http import HttpResponse
from django.utils.html import escape
from django.utils.safestring import mark_safe
from django.utils.translation import gettext_lazy as _
from django_json_widget.widgets import JSONEditorWidget
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from parler.admin import TranslatableAdmin

from apps import models
from apps.models import (
    User, Country, Question, SalesOutlets, Place, Category,
    Event, Ticket, Wishlist, OrderItem, Order, Payment, Transaction, Address
)



admin.sites.site.unregister(Group)


@admin.action(description=_("Export selected objects to Excel (.xlsx)"))
def export_to_excel(modeladmin, request, queryset):
    """
    Generic action to export any model data to Excel (.xlsx) format.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = modeladmin.model._meta.verbose_name_plural.capitalize()[:31]

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    fields = [field for field in modeladmin.model._meta.fields]
    headers = [field.verbose_name.title() for field in fields]

    ws.append(headers)
    ws.row_dimensions[1].height = 25

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, obj in enumerate(queryset, 2):
        row_data = []
        for field in fields:
            val = getattr(obj, field.name)
            if isinstance(val, (datetime.datetime, datetime.date)):
                val = val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(val, datetime.datetime) else val.strftime(
                    "%Y-%m-%d")
            elif hasattr(val, '__str__') and not isinstance(val, (int, float, str, bool, type(None))):
                val = str(val)
            row_data.append(val if val is not None else "")

        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20

        for col_num in range(1, len(fields) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"{modeladmin.model._meta.model_name}_export_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@admin.register(User)
class UserAdmin(BaseUserAdmin):
    list_display = ("id", "phone_number", "email", "gender", "birth_date", "country", "is_staff", "is_active",
                    "date_joined")
    list_filter = ("gender", "is_staff", "is_active", "is_superuser", "country")
    search_fields = ("phone_number", "email", "first_name", "last_name")
    ordering = ("-id",)
    list_per_page = 25
    filter_horizontal = ("user_permissions",)
    actions = [export_to_excel]

    fieldsets = (
        (None, {"fields": ("phone_number", "password")}),
        (_("Personal Info"), {"fields": ("first_name", "last_name", "email", "gender", "birth_date", "country")}),
        (_("Permissions"), {"fields": ("is_active", "is_staff", "is_superuser", "user_permissions")}),
        (_("Important dates"), {"fields": ("last_login", "date_joined")}),
    )
    add_fieldsets = (
        (
            None,
            {
                "classes": ("wide",),
                "fields": ("phone_number", "email", "password1", "password2"),
            },
        ),
    )


@admin.register(Country)
class CountryAdmin(TranslatableAdmin):
    list_display = ("id", "name")
    search_fields = ("translations__name",)
    list_per_page = 25
    actions = [export_to_excel]


@admin.register(Question)
class QuestionAdmin(admin.ModelAdmin):
    list_display = ("id", "short_question", "short_answer", "is_visible")
    search_fields = ("question", "answer")
    list_per_page = 25
    actions = [export_to_excel]

    @admin.display(description=_("Question"))
    def short_question(self, obj):
        return obj.question[:60] + "..." if len(obj.question) > 60 else obj.question

    @admin.display(description=_("Answer"))
    def short_answer(self, obj):
        return obj.answer[:60] + "..." if len(obj.answer) > 60 else obj.answer


@admin.register(SalesOutlets)
class SalesOutletsAdmin(admin.ModelAdmin):
    list_display = ("id", "place", "start_time", "end_time", "latitude", "longitude")
    list_filter = ("place",)
    search_fields = ("place__translations__title", "place__translations__name")
    list_per_page = 25
    actions = [export_to_excel]
    formfield_overrides = {
        models.JSONField: {'widget': JSONEditorWidget},
    }


@admin.register(Place)
class PlaceAdmin(TranslatableAdmin):
    list_display = ("id", "phone_number", "image_preview")
    search_fields = ("phone_number", "translations__title", "translations__name")
    list_per_page = 25
    actions = [export_to_excel]

    @admin.display(description=_("Image"))
    def image_preview(self, obj):
        if obj.image:
            return mark_safe(
                f'<img src="{escape(obj.image.url)}" style="width: 45px; height:45px; object-fit:cover; border-radius:6px;" />')
        return "-"


@admin.register(Category)
class CategoryAdmin(TranslatableAdmin):
    list_display = ("id", "name")
    search_fields = ("translations__name",)
    list_per_page = 25
    actions = [export_to_excel]


class TicketInline(admin.TabularInline):
    model = Ticket
    extra = 1


@admin.register(Event)
class EventAdmin(TranslatableAdmin):
    list_display = ("id", "category", "place", "start_datetime", "end_datetime", "image_preview")
    list_filter = ("category", "place", "start_datetime")
    inlines = [TicketInline]
    list_per_page = 25
    actions = [export_to_excel]
    formfield_overrides = {
        models.JSONField: {'widget': JSONEditorWidget},
    }

    @admin.display(description=_("Image"))
    def image_preview(self, obj):
        if obj.image:
            return mark_safe(
                f'<img src="{escape(obj.image.url)}" style="width: 45px; height:45px; object-fit:cover; border-radius:6px;" />')
        return "-"


@admin.register(Ticket)
class TicketAdmin(TranslatableAdmin):
    list_display = ("id", "price", "count", "even")
    list_filter = ("even",)
    search_fields = ("translations__title", "even__translations__title")
    list_per_page = 25
    actions = [export_to_excel]


@admin.register(Wishlist)
class WishlistAdmin(admin.ModelAdmin):
    list_display = ("id", "user", "event", "created_at", "updated_at")
    list_filter = ("created_at",)
    search_fields = ("user__phone_number", "user__email")
    readonly_fields = ("created_at", "updated_at")
    list_per_page = 25
    actions = [export_to_excel]


@admin.register(OrderItem)
class OrderItemAdmin(admin.ModelAdmin):
    list_display = ("id", "ticket", "order", "count", "updated_at", "created_at")
    list_filter = ("created_at", "updated_at")
    search_fields = ("user__phone_number", "user__email", "ticket__translations__title")
    readonly_fields = ("created_at", "updated_at")
    list_per_page = 25
    actions = [export_to_excel]


class PaymentInline(admin.TabularInline):
    model = Payment
    extra = 0
    readonly_fields = ("created_at", "updated_at")


@admin.register(Order)
class OrderAdmin(admin.ModelAdmin):
    list_display = ("id", "user", "total_amount", "colored_status", "created_at", "updated_at")
    list_filter = ("status", "created_at")
    search_fields = ("user__phone_number", "user__email", "id")
    readonly_fields = ("created_at", "updated_at")
    inlines = [PaymentInline]
    list_per_page = 25
    actions = [export_to_excel, "mark_as_delivered", "mark_as_cancelled"]

    @admin.display(description=_("Status"))
    def colored_status(self, obj):
        colors = {
            "pending": "#f39c12",
            "delivered": "#2ecc71",
            "cancelled": "#e74c3c",
            "failed": "#95a5a6",
        }
        color = colors.get(obj.status, "#333")
        return mark_safe(
            f'<span style="background-color: {escape(color)}; color: white; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 11px;">{escape(obj.get_status_display())}</span>'
        )

    @admin.action(description=_("Mark selected orders as Delivered"))
    def mark_as_delivered(self, request, queryset):
        updated = queryset.update(status=Order.StatusTextChoices.DELIVERED)
        self.message_user(request, f"{updated} order(s) marked as Delivered.", messages.SUCCESS)

    @admin.action(description=_("Mark selected orders as Cancelled"))
    def mark_as_cancelled(self, request, queryset):
        updated = queryset.update(status=Order.StatusTextChoices.CANCELLED)
        self.message_user(request, f"{updated} order(s) marked as Cancelled.", messages.WARNING)


@admin.register(Payment)
class PaymentAdmin(admin.ModelAdmin):
    list_display = ("id", "order", "total_amount", "colored_status", "created_at", "updated_at")
    list_filter = ("status", "created_at")
    search_fields = ("order__id", "total_amount")
    readonly_fields = ("created_at", "updated_at")
    list_per_page = 25
    actions = [export_to_excel, "mark_as_completed", "mark_as_refunded"]

    @admin.display(description=_("Status"))
    def colored_status(self, obj):
        colors = {
            "pending": "#f39c12",
            "completed": "#2ecc71",
            "cancelled": "#e74c3c",
            "refunded": "#9b59b6",
        }
        color = colors.get(obj.status, "#333")
        return mark_safe(
            f'<span style="background-color: {escape(color)}; color: white; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 11px;">{escape(obj.get_status_display())}</span>'
        )

    @admin.action(description=_("Mark selected payments as Completed"))
    def mark_as_completed(self, request, queryset):
        updated = queryset.update(status=Payment.StatusTextChoices.COMPLETED)
        self.message_user(request, f"{updated} payment(s) marked as Completed.", messages.SUCCESS)

    @admin.action(description=_("Mark selected payments as Refunded"))
    def mark_as_refunded(self, request, queryset):
        updated = queryset.update(status=Payment.StatusTextChoices.REFUNDED)
        self.message_user(request, f"{updated} payment(s) marked as Refunded.", messages.INFO)


@admin.register(Transaction)
class TransactionAdmin(admin.ModelAdmin):
    list_display = ("id", "payment", "colored_status", "created_at", "updated_at")
    list_filter = ("status", "created_at")
    search_fields = ("payment__id",)
    readonly_fields = ("created_at", "updated_at")
    list_per_page = 25
    actions = [export_to_excel]

    @admin.display(description=_("Status"))
    def colored_status(self, obj):
        colors = {
            "success": "#2ecc71",
            "failed": "#e74c3c",
        }
        color = colors.get(obj.status, "#333")
        return mark_safe(
            f'<span style="background-color: {escape(color)}; color: white; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 11px;">{escape(obj.get_status_display())}</span>'
        )


@admin.register(Address)
class AddressAdmin(admin.ModelAdmin):
    list_display = ("id", "title", "city", "street", "building", "user", "country")
    list_filter = ("city", "country")
    search_fields = ("title", "city", "street", "user__phone_number", "user__email")
    list_per_page = 25
    actions = [export_to_excel]
